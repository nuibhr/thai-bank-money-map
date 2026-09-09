from __future__ import annotations
import csv
import json
import re
from pathlib import Path
from typing import Iterator

import fitz
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation

from .chunking import chunk_text, stable_id, clean_text
from .config import BANKS

SUPPORTED = {".pdf", ".xlsx", ".xlsm", ".csv", ".tsv", ".docx", ".pptx", ".txt", ".md", ".json", ".html", ".htm"}


def infer_ticker(path: Path) -> str | None:
    hay = " ".join(path.parts + (path.stem,)).upper()
    for t in BANKS:
        if re.search(rf"\b{re.escape(t)}\b", hay):
            return t
    return None


def infer_doc_type(path: Path) -> str:
    n = path.name.lower()
    rules = [
        ("pillar", "pillar3"), ("mda", "mda"), ("management discussion", "mda"),
        ("analyst", "analyst_presentation"), ("presentation", "presentation"),
        ("annual", "annual_report"), ("financial", "financial_statement"),
        ("fs", "financial_statement"), ("results", "results"),
    ]
    for key, val in rules:
        if key in n:
            return val
    return path.suffix.lower().lstrip(".") or "document"


def _emit(path: Path, unit: str, text: str, extra: dict | None = None) -> Iterator[dict]:
    text = clean_text(text)
    if not text:
        return
    sidecar = path.with_suffix(path.suffix + ".meta.json")
    supplied = json.loads(sidecar.read_text(encoding="utf-8")) if sidecar.exists() else {}
    ticker = supplied.get("ticker") or infer_ticker(path)
    if not ticker or not supplied.get("period"):
        raise ValueError("Document requires ticker and period in .meta.json sidecar")
    doc_type = infer_doc_type(path)
    for i, chunk in enumerate(chunk_text(text)):
        meta = {
            "path": str(path),
            "filename": path.name,
            "ticker": ticker,
            "doc_type": doc_type,
            "unit": unit,
            "chunk_index": i,
            "source_kind": "local_document",
        }
        meta.update({k: supplied[k] for k in ("period", "sourceUrl", "publishedAt", "entityScope") if k in supplied})
        if extra:
            meta.update(extra)
        yield {"id": stable_id(path, unit, i, chunk[:80]), "text": chunk, "metadata": meta}


def load_pdf(path: Path) -> Iterator[dict]:
    doc = fitz.open(path)
    for pno, page in enumerate(doc, start=1):
        text = page.get_text("text") or ""
        yield from _emit(path, f"page:{pno}", text, {"page": pno})


def load_xlsx(path: Path) -> Iterator[dict]:
    wb = load_workbook(path, read_only=True, data_only=False)
    for ws in wb.worksheets:
        batch: list[str] = []
        start_row = 1
        for ridx, row in enumerate(ws.iter_rows(values_only=True), start=1):
            vals = ["" if v is None else str(v) for v in row]
            if not any(v.strip() for v in vals):
                continue
            line = " | ".join(vals)
            if not batch:
                start_row = ridx
            batch.append(line)
            if len(batch) >= 30:
                yield from _emit(path, f"sheet:{ws.title}!{start_row}:{ridx}", "\n".join(batch), {"sheet": ws.title, "row_start": start_row, "row_end": ridx})
                batch = []
        if batch:
            end_row = start_row + len(batch) - 1
            yield from _emit(path, f"sheet:{ws.title}!{start_row}:{end_row}", "\n".join(batch), {"sheet": ws.title, "row_start": start_row, "row_end": end_row})


def load_csv_like(path: Path) -> Iterator[dict]:
    delim = "\t" if path.suffix.lower() == ".tsv" else ","
    with path.open("r", encoding="utf-8-sig", errors="replace", newline="") as f:
        reader = csv.reader(f, delimiter=delim)
        batch, start = [], 1
        for ridx, row in enumerate(reader, start=1):
            if not batch:
                start = ridx
            batch.append(" | ".join(row))
            if len(batch) >= 50:
                yield from _emit(path, f"rows:{start}:{ridx}", "\n".join(batch), {"row_start": start, "row_end": ridx})
                batch = []
        if batch:
            yield from _emit(path, f"rows:{start}:{start+len(batch)-1}", "\n".join(batch), {"row_start": start, "row_end": start+len(batch)-1})


def load_docx(path: Path) -> Iterator[dict]:
    doc = Document(path)
    paras = [p.text for p in doc.paragraphs if p.text.strip()]
    yield from _emit(path, "document", "\n\n".join(paras))


def load_pptx(path: Path) -> Iterator[dict]:
    prs = Presentation(path)
    for i, slide in enumerate(prs.slides, start=1):
        texts = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                texts.append(shape.text)
        yield from _emit(path, f"slide:{i}", "\n".join(texts), {"slide": i})


def load_json(path: Path) -> Iterator[dict]:
    with path.open("r", encoding="utf-8", errors="replace") as f:
        obj = json.load(f)
    text = json.dumps(obj, ensure_ascii=False, indent=2)
    yield from _emit(path, "json", text)


def load_html(path: Path) -> Iterator[dict]:
    from bs4 import BeautifulSoup
    html = path.read_text(encoding="utf-8", errors="replace")
    soup = BeautifulSoup(html, "html.parser")
    yield from _emit(path, "html", soup.get_text("\n"))


def load_text(path: Path) -> Iterator[dict]:
    yield from _emit(path, "document", path.read_text(encoding="utf-8", errors="replace"))


def load_file(path: Path) -> Iterator[dict]:
    ext = path.suffix.lower()
    if ext not in SUPPORTED:
        return iter(())
    if ext == ".pdf": return load_pdf(path)
    if ext in {".xlsx", ".xlsm"}: return load_xlsx(path)
    if ext in {".csv", ".tsv"}: return load_csv_like(path)
    if ext == ".docx": return load_docx(path)
    if ext == ".pptx": return load_pptx(path)
    if ext == ".json": return load_json(path)
    if ext in {".html", ".htm"}: return load_html(path)
    return load_text(path)
